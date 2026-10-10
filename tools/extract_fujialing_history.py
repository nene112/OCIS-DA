"""Read the calibration workbook without modifying it; retain row provenance."""
import argparse
import hashlib
import json
from pathlib import Path
import openpyxl
import pandas as pd
import numpy as np

ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output',type=Path,default=ROOT/'data/dayudu/output/fujialing_reconstruction_comparison')
args=parser.parse_args();out=args.output;out.mkdir(parents=True,exist_ok=True)
source=ROOT/'data/dayudu/input/闸门水位流量详情-富家岭-模型率定用.xlsx'
sheet=openpyxl.load_workbook(source,data_only=True,read_only=True).active
headers=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True))
rows=[]
for index,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
    if row[4] is None:continue
    rows.append({'source_row':index,'time':row[4],'gate_name':row[1],
        'station2_inflow_m3s':row[0],'h_down_record_m':row[2],'h_up_record_m':row[3],
        'opening_1_raw_mm':row[5],'opening_2_raw_mm':row[6],'q_reference_calculated_m3s':row[7]})
data=pd.DataFrame(rows)
numeric=[c for c in data if c not in ['time','gate_name']]
for c in numeric:data[c]=pd.to_numeric(data[c],errors='coerce')
data['time']=pd.to_datetime(data.time)
data['opening_1_m']=data.opening_1_raw_mm/1000
data['opening_2_m']=data.opening_2_raw_mm/1000
data['opening_sum_m']=data.opening_1_m+data.opening_2_m
data['head_difference_record_m']=data.h_up_record_m-data.h_down_record_m
data['valid_finite']=np.isfinite(data[numeric]).all(axis=1)
data['valid_opening']=(data.opening_1_raw_mm>=0)&(data.opening_2_raw_mm>=0)
data['valid_forward_head']=(data.h_up_record_m>0)&(data.h_down_record_m>=0)&(data.head_difference_record_m>=0)
data['valid_q']=data.q_reference_calculated_m3s>=0
data['relation_valid']=data[['valid_finite','valid_opening','valid_forward_head','valid_q']].all(axis=1)
data=data.sort_values(['time','source_row'])
data.to_csv(out/'historical_records_all.csv',index=False,encoding='utf-8-sig')
data[data.relation_valid].to_csv(out/'historical_relation_valid.csv',index=False,encoding='utf-8-sig')
report={'source':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
    'sheet':sheet.title,'source_range':'A1:H3777','headers':list(headers[:8]),
    'records':len(data),'valid_relation_records':int(data.relation_valid.sum()),
    'negative_opening_records_excluded':int((~data.valid_opening).sum()),
    'invalid_head_records_excluded':int((~data.valid_forward_head).sum()),
    'period':[str(data.time.min()),str(data.time.max())],
    'opening_1_max_m':float(data.opening_1_m.max()),'opening_2_max_m':float(data.opening_2_m.max()),
    'reference_flow_is_calculated_not_independent_measurement':True,
    'station2_inflow_is_not_fujialing_gate_discharge':True,
    'negative_opening_policy':'exclude; do not silently clamp sensor -2 mm to closed',
    'water_record_interpretation':'values compared to model depth h (roughly 0-3 m); datum relative to gate sill is not documented'}
(out/'historical_data_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
